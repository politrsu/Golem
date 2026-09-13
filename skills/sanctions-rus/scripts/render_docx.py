#!/usr/bin/env python3
import json
import re
import sys
from pathlib import Path
from docx import Document

CATEGORY_TITLES = {
    "individuals": "Физические лица",
    "entities": "Юридические лица и организации",
    "vessels": "Суда",
    "aircraft": "Воздушные суда",
    "banks": "Банки и финансовые организации",
    "ports": "Порты",
}
UKRAINIAN = re.compile(r"[іїєґІЇЄҐ]")


def label(item):
    name = item.get("name_ru") or item.get("name")
    if not name:
        raise ValueError("Each record requires name_ru")
    extras = []
    if item.get("position_ru"):
        extras.append(item["position_ru"])
    if item.get("imo"):
        extras.append(f"IMO {item['imo']}")
    if item.get("country_ru"):
        extras.append(item["country_ru"])
    return name + (" — " + "; ".join(extras) if extras else "")


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: render_docx.py EVENT.json OUTPUT.docx")
    event = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    raw = json.dumps(event, ensure_ascii=False)
    if UKRAINIAN.search(raw):
        raise SystemExit("Ukrainian-specific letters remain; translate all content into Russian first")
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.add_heading(event.get("title_ru", "Санкционный список"), 0)
    for key, items in event.get("categories", {}).items():
        if not items:
            continue
        document.add_heading(CATEGORY_TITLES.get(key, key), level=1)
        for item in items:
            document.add_paragraph(label(item), style="List Bullet")
    if event.get("official_url"):
        document.add_paragraph(f"Оригинал документа: {event['official_url']}")
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    main()
