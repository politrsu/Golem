"""Synthetic, offline checks of Word structure and document content gates."""
import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location('render_reference', ROOT / 'scripts/render_docx.py')
RENDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(RENDER)


def event():
    return {
        'title_ru': 'Пример: новые включения', 'published_date': '2030-01-02',
        'official_url': 'https://authority.example/notice',
        'legal_effect_ru': 'Дополнительно изменены ограничения для ранее включённых организаций.',
        'categories': {
            'individuals': [{'name_ru': 'Иван Примеров', 'country_ru': 'Казахстан',
                'position_ru': 'Совладелец ООО «Пример»; управленческая должность не установлена',
                'position_evidence': [{'source_url': 'https://authority.example/person'}]}],
            'entities': [{'name_ru': 'ООО «Пример»', 'country_ru': 'Россия'},
                         {'name_ru': 'Example Ltd', 'country_ru': 'Китай'}],
            'vessels': [{'name_ru': 'Example Vessel', 'imo': 'IMO 1234567'}], 'aircraft': []}}


class ReferenceDocxTests(unittest.TestCase):
    def render(self, data):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'list.docx'
            RENDER.render(data, output)
            return Document(output)

    def test_layout_numbering_links_roles_and_foreign_countries(self):
        doc = self.render(event())
        self.assertEqual(doc.styles['Normal'].font.name, 'Verdana')
        self.assertEqual(doc.styles['Normal'].font.size.pt, 11)
        self.assertAlmostEqual(doc.styles['Normal'].paragraph_format.line_spacing, 1.15)
        for key, value in {'page_width': 210, 'page_height': 297, 'top_margin': 20,
                           'bottom_margin': 20, 'left_margin': 30, 'right_margin': 15}.items():
            self.assertAlmostEqual(getattr(doc.sections[0], key).mm, value, delta=0.03)
        self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertIn('Новых позиций: 4', doc.paragraphs[1].text)
        self.assertTrue(doc.paragraphs[1].runs[0].italic)
        self.assertFalse(doc.tables)
        self.assertEqual(doc.core_properties.author, '')
        self.assertEqual(doc.core_properties.last_modified_by, '')
        numbered = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
        self.assertEqual(len(numbered), 4)
        self.assertIn('Иван Примеров (Казахстан) – Совладелец', numbered[0].text)
        self.assertTrue(numbered[0].runs[0].bold)
        self.assertNotIn('(Россия)', numbered[1].text)
        self.assertIn('(Китай)', numbered[2].text)
        self.assertEqual(numbered[3].text.count('IMO '), 1)
        ids = [p._p.pPr.numPr.numId.val for p in numbered]
        self.assertEqual(ids[1], ids[2])
        self.assertEqual(len(set(ids)), 3)
        root = doc.part.numbering_part.element
        for number in set(ids):
            num = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == str(number))
            aid = num.find(qn('w:abstractNumId')).get(qn('w:val'))
            abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == aid)
            level = abstract.find(qn('w:lvl'))
            self.assertEqual(level.find(qn('w:start')).get(qn('w:val')), '1')
            self.assertEqual(level.find(qn('w:numFmt')).get(qn('w:val')), 'decimal')
        links = {r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK}
        self.assertEqual(links, {'https://authority.example/notice', 'https://authority.example/person'})
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertNotIn('Воздушные суда', text)
        self.assertIn(event()['legal_effect_ru'], text)

    def test_missing_or_placeholder_role_preserves_existing_artifact(self):
        for role in (None, '', '   ', 'должность', 'Должность не установлена', 'нет данных', 'N/A', 'unknown', '—'):
            with self.subTest(role=role), tempfile.TemporaryDirectory() as directory:
                data = event()
                data['categories']['individuals'][0]['position_ru'] = role
                output = Path(directory) / 'existing.docx'
                output.write_bytes(b'previous reviewed artifact')
                with self.assertRaises(ValueError):
                    RENDER.render(data, output)
                self.assertEqual(output.read_bytes(), b'previous reviewed artifact')

    def test_missing_or_invalid_role_evidence_blocks_render(self):
        for evidence in (None, [], [{}], [{'source_url': 'file:///private'}], [{'source_url': 'javascript:alert(1)'}]):
            with self.subTest(evidence=evidence):
                data = event()
                data['categories']['individuals'][0]['position_evidence'] = evidence
                with self.assertRaises(ValueError):
                    self.render(data)

    def test_original_ukrainian_evidence_is_preserved_but_not_rendered(self):
        data = event()
        person = data['categories']['individuals'][0]
        person['name'] = 'Оригінальне ім’я'
        person['position_evidence'][0]['quote_original'] = 'Керівник підприємства'
        original = copy.deepcopy(data)
        doc = self.render(data)
        self.assertEqual(data, original)
        self.assertNotIn('Керівник', '\n'.join(p.text for p in doc.paragraphs))
        person['position_ru'] = 'Керівник підприємства'
        with self.assertRaises(ValueError):
            self.render(data)

    def test_unknown_category_and_bad_imo_are_rejected(self):
        data = event()
        data['categories']['unknown'] = [{'name_ru': 'Пример'}]
        with self.assertRaises(ValueError):
            self.render(data)
        data = event()
        data['categories']['vessels'][0]['imo'] = '123'
        with self.assertRaises(ValueError):
            self.render(data)


if __name__ == '__main__':
    unittest.main()
